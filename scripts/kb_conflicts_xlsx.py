"""Противоречия базы знаний → Excel владельцу на разбор (30.10): правила > PREMIUM MARKET > GameNonStop.

Читает data/kb/rule_conflicts.json (services.dialog.merge_rules) и data/kb/conflicts_<класс>.json
(services.dialog.build_kb). Кандидатов LLM находит сотни, большая часть — одно и то же разными словами или одна тема
в нескольких разделах (30.10: 565 строк), поэтому LLM ещё раз отбирает настоящие споры и склеивает повторы.
Лист «Противоречия»: строка — спорный вопрос, колонки — что говорит каждый источник по порядку и что сейчас в базе;
сначала «PREMIUM ≠ GameNonStop» (правил нет — решает владелец), потом «Правило ≠ переписка» (действует правило).
Лист «Правила между собой» — споры правил владельца. Верное выделяют цветом или пишут в «Как надо».
Запуск (на ПК): python scripts/kb_conflicts_xlsx.py [каталог с json] [файл.xlsx]
"""
import asyncio
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from services.dialog.build_kb import NOTHING, SIDES, STEAM  # noqa: E402
from shared.llm import LLM  # noqa: E402

PM_GNS, RULE = "PREMIUM ≠ GameNonStop", "Правило ≠ переписка"
HEAD = ("№", "Тип", "Раздел", "Вопрос / ситуация", "Правила", "PREMIUM MARKET", "GameNonStop", "Сейчас в базе",
        "Как надо (если верного нет)")
KEYS = ("kind", "section", "topic", "rules", "premium", "gamenonstop", "chosen")
WIDTH = (5, 14, 22, 30, 40, 40, 40, 40, 30)
BATCH = 40

REVIEW = """Ниже — кандидаты в противоречия базы знаний бота-продавца подписки Xbox Game Pass (Avito).
R — правила бота и владельца, P — переписка магазина PREMIUM MARKET, G — переписка магазина GameNonStop.
Для каждого реши, настоящее ли это противоречие: источники по сути велят ответить клиенту или поступить по-разному —
разные факты или условия, «да» против «нет», «передать человеку» против «ответить самим», разные шаги.
Не противоречие: одно и то же разными словами; один источник подробнее или уже другого; разные стороны одного
вопроса; один источник пустой. Каждому настоящему дай ключ темы — 3–6 слов о сути вопроса клиента.
Ответ — JSON: {"items": [{"n": номер, "real": true или false, "key": "ключ темы"}]}"""
GROUP = """Ниже — ключи тем спорных вопросов базы знаний бота (подписка Xbox Game Pass, Avito), по одному в строке.
Объедини ключи об одном и том же вопросе клиента. Ответ — JSON: {"groups": [["ключ", "ключ", ...], ...]};
каждый ключ — ровно в одной группе, без пары — группа из одного."""


def load(src: Path) -> tuple[list[dict], list[dict]]:
    """(споры правил между собой, кандидаты из базы) — в кандидатах «Пусто»/«не упоминается» уже пустые."""
    rules = json.loads(p.read_text(encoding="utf-8")) if (p := src / "rule_conflicts.json").exists() else []
    rows = []
    for f in sorted(src.glob("conflicts_*.json")):
        for c in json.loads(f.read_text(encoding="utf-8")):
            c = {**c, **{k: "" for k in SIDES if NOTHING.match(c.get(k) or "")}}
            if STEAM.search(" ".join(c.get(k) or "" for k in ("topic", *SIDES))):
                continue   # Steam не продаём (30.10) — спорить не о чем, из базы он вычищен
            if sum(bool(c.get(k)) for k in SIDES) >= 2:
                rows.append({**c, "kind": PM_GNS if not c["rules"] else RULE})
    return rules, rows


def line(n: int, c: dict) -> str:
    return f"{n}. «{c['topic']}» | R: {c['rules'] or '—'} | P: {c['premium'] or '—'} | G: {c['gamenonstop'] or '—'}"


async def review(llm: LLM, rows: list[dict]) -> list[dict]:
    """Только настоящие споры, у каждого ключ темы (rows[i]["key"])."""
    async def part(start: int) -> list[dict]:
        chunk = rows[start:start + BATCH]
        res = await llm.chat_json(REVIEW, "\n".join(line(i, c) for i, c in enumerate(chunk, 1)),
                                  max_tokens=4000, temperature=0)
        out = []
        for it in res.get("items") or []:
            n = int(it["n"]) if isinstance(it, dict) and str(it.get("n")).isdigit() else 0
            if it.get("real") is True and 1 <= n <= len(chunk):
                out.append({**chunk[n - 1], "key": " ".join(str(it.get("key") or chunk[n - 1]["topic"]).split())})
        return out

    return [c for p in await asyncio.gather(*(part(i) for i in range(0, len(rows), BATCH))) for c in p]


def merge_groups(rows: list[dict], groups: list[list[str]]) -> list[dict]:
    """Одна строка на тему: самая полная (больше источников, длиннее), разделы — все, где встретилась."""
    of = {k: i for i, g in enumerate(groups) for k in g}
    by: dict = {}
    for c in rows:
        by.setdefault((of.get(c["key"], c["key"]), c["kind"]), []).append(c)
    out = []
    for same in by.values():
        best = max(same, key=lambda c: (sum(bool(c[k]) for k in SIDES), sum(len(c[k]) for k in SIDES)))
        out.append({**best, "section": "; ".join(dict.fromkeys(c["section"] for c in same))})
    return sorted(out, key=lambda c: (c["kind"] != PM_GNS, c["section"]))


async def select(llm: LLM, rows: list[dict]) -> list[dict]:
    real = await review(llm, rows)
    keys = list(dict.fromkeys(c["key"] for c in real))
    groups = (await llm.chat_json(GROUP, "\n".join(keys), max_tokens=8000, temperature=0)).get("groups") or [] \
        if len(keys) > 1 else []
    return merge_groups(real, [g for g in groups if isinstance(g, list)])


def sheet(ws, head, rows, width) -> None:
    ws.append(head)
    for r in rows:
        ws.append(r)
    for col, w in enumerate(width, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    ws.freeze_panes = "E2" if len(head) > 5 else "C2"
    ws.auto_filter.ref = ws.dimensions


def build(rules: list[dict], rows: list[dict], out: Path) -> None:
    wb = Workbook()
    sheet(wb.active, HEAD, [[i, *(c.get(k) or "" for k in KEYS), ""] for i, c in enumerate(rows, 1)], WIDTH)
    wb.active.title = "Противоречия"
    if rules:
        sheet(wb.create_sheet("Правила между собой"), ("№", "Ситуация", "Действует", "Выключено", "Как надо"),
              [[i, c.get("topic") or "", c.get("chosen") or "", "\n".join(c.get("off") or []), ""]
               for i, c in enumerate(rules, 1)], (5, 30, 60, 60, 30))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


async def main(src: Path, out: Path) -> None:
    rules, rows = load(src)
    llm = LLM()
    try:
        picked = await select(llm, rows)
    finally:
        await llm.close()
    build(rules, picked, out)
    print(f"кандидатов {len(rows)} → споров {len(picked)} "
          f"({sum(c['kind'] == PM_GNS for c in picked)} PREMIUM ≠ GameNonStop), правил между собой {len(rules)} → {out}")


if __name__ == "__main__":
    asyncio.run(main(Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "kb",
                     Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "docs" / "База знаний — противоречия.xlsx"))
