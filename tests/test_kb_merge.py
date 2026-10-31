"""База знаний одна на оба аккаунта (30.10): правила > PREMIUM MARKET > GameNonStop, противоречия — владельцу."""
import asyncio
import importlib.util
import json
from pathlib import Path

import pytest

from services.dialog import build_kb
from services.dialog.merge_rules import winner

ROOT = Path(__file__).resolve().parent.parent


def test_rule_winner_premium_then_fresh():
    tags = {1: "G", 2: "G", 3: "P", 4: "PG"}
    assert winner([1, 3], tags) == 3                # PREMIUM важнее GameNonStop, даже старее по id
    assert winner([1, 2], tags) == 2                # один аккаунт — свежее
    assert winner([3, 4], tags) == 4                # оба PREMIUM — свежее
    assert winner([1, 2], tags, prefer=1) == 1      # повтор: формулировка, которую выбрала LLM
    assert winner([1, 3], tags, prefer=1) == 3      # …но не в ущерб PREMIUM


def test_clean_conflicts_needs_two_sides():
    out = build_kb.clean_conflicts([
        {"topic": "PS Plus", "rules": "Уточню и напишу", "premium": "", "gamenonstop": "Да, есть",
         "chosen": "Уточню"},
        {"topic": "только одно", "premium": "что-то", "gamenonstop": " "},
        "мусор", {"topic": "x", "premium": ["список"], "gamenonstop": 5}], "Game Pass · Частые вопросы")
    assert [c["topic"] for c in out] == ["PS Plus", "x"]
    assert out[0]["section"] == "Game Pass · Частые вопросы" and out[0]["chosen"] == "Уточню"


def chat(chat_id, *texts):
    return {"chat_id": chat_id, "messages": [{"role": "seller", "text": t} for t in texts]}


def test_templates_premium_first():
    pm, gns = build_kb.Source("PREMIUM MARKET", "2"), build_kb.Source("GameNonStop", "")
    hello = "Здравствуйте! Подключаем Game Pass Ultimate на Ваш аккаунт"
    pm.by_class = {"core": [chat(f"p{i}", hello, "Код придёт на почту, пришлите его сюда") for i in range(3)]}
    gns.by_class = {"core": [chat(f"g{i}", "Доброго дня, подписка работает на консоли и ПК", hello)
                             for i in range(5)]}
    out = [t for t, _ in build_kb.templates_by_priority([pm, gns], "core")]
    assert out[:2] == [hello, "Код придёт на почту, пришлите его сюда"] or \
        out[:2] == ["Код придёт на почту, пришлите его сюда", hello]
    assert out[2] == "Доброго дня, подписка работает на консоли и ПК" and len(out) == 3
    assert gns.export == ROOT / "data" / "export" and pm.work == ROOT / "data2" / "kb"


class FakeLLM:
    def __init__(self):
        self.calls = []

    async def chat(self, messages, json_mode=False, **kw):
        self.calls.append(messages[1]["content"])
        if json_mode:
            return ('{"conflicts": [{"topic": "чек", "rules": "не выдаём", "premium": "", '
                    '"gamenonstop": "по запросу", "chosen": "не выдаём"}]}')
        return "Подписка PREMIUM MARKET работает на консоли"


def test_prioritize_order_cache_and_store_names(tmp_path, monkeypatch):
    monkeypatch.setattr(build_kb, "WORK", tmp_path)
    parts = [("PREMIUM MARKET", {k: f"P-{k}" for k in build_kb.REDUCE}),
             ("GameNonStop", {k: f"G-{k}" for k in build_kb.REDUCE})]
    llm = FakeLLM()
    bodies, conflicts = asyncio.run(build_kb.prioritize(llm, "core", "ПРАВИЛА-X", parts))
    data = llm.calls[0]
    assert data.index("ПРАВИЛА-X") < data.index("### PREMIUM MARKET") < data.index("### GameNonStop")
    # «почему не купили» — аналитика: сводится, но противоречия там не ищем
    assert len(llm.calls) == 2 * len(build_kb.REDUCE) - 1 and len(conflicts) == len(build_kb.REDUCE) - 1
    assert conflicts[0]["section"].endswith("Товар и условия")
    n = len(llm.calls)
    asyncio.run(build_kb.prioritize(llm, "core", "ПРАВИЛА-X", parts))
    assert len(llm.calls) == n                        # ничего не изменилось — из кэша
    asyncio.run(build_kb.prioritize(llm, "core", "ПРАВИЛА-Y", parts))
    assert len(llm.calls) > n                         # новое правило — сводим заново
    pm = build_kb.Source("PREMIUM MARKET", "2", by_class={"core": []})
    text = build_kb.assemble("core", bodies, *build_kb.stats("core", [pm]), "прайс")
    assert "PREMIUM MARKET работает" not in text and "наш магазин работает" in text   # база общая на оба бота
    assert "правила > PREMIUM MARKET > GameNonStop" in text


class ScriptLLM:
    """Сведение раздела отвечает по списку; поиск противоречий — пусто."""
    def __init__(self, answers):
        self.answers, self.calls = answers, 0

    async def chat(self, messages, json_mode=False, **kw):
        if json_mode:
            return '{"conflicts": []}'
        self.calls += 1
        return self.answers.pop(0)


def test_prioritize_rejects_bot_reply_and_thin(tmp_path, monkeypatch):
    """30.10 прод: вместо раздела «Возражения» — JSON-ответ клиенту, раздел игр потерял 4/5 фактов."""
    monkeypatch.setattr(build_kb, "WORK", tmp_path)
    monkeypatch.setattr(build_kb, "REDUCE", {"objections": ("Возражения", "Сведи возражения.")})
    parts = [("PREMIUM MARKET", {"objections": "П" * 100}), ("GameNonStop", {"objections": "Г" * 300})]
    run = lambda llm, rules: asyncio.run(build_kb.prioritize(llm, "core", rules, parts))[0]["objections"]
    llm = ScriptLLM(['```json\n{"reply": "Здравствуйте! На какой срок?", "handoff": null}\n```', "сведено " * 40])
    assert run(llm, "R1") == "сведено " * 40 and llm.calls == 2
    llm = ScriptLLM(['{"reply": "а"}', '{ "reply" : "б"}'])
    assert run(llm, "R2") == "П" * 100                      # обе попытки — ответ клиенту: самый важный источник
    llm = ScriptLLM(["коротко", "чуть длиннее"])
    assert run(llm, "R3") == "чуть длиннее"                 # обе тонкие — самая полная из попыток


def test_rules_text_without_bot_format():
    text = build_kb.rules_text(["Чек не выдаём"])
    assert "ФОРМАТ ОТВЕТА" not in text and '"reply"' not in text and text.startswith("ПРАВИЛА БОТА (справка")
    assert text.endswith("ПРАВИЛА ВЛАДЕЛЬЦА:\n- Чек не выдаём")


class ReviewLLM:
    """Отбор: «чек» — настоящий спор (в двух разделах), «оплата» — одно и то же разными словами."""
    async def chat_json(self, system, user, **kw):
        if "ключи тем" in system:
            return {"groups": [["чек клиенту", "выдаём ли чек"]]}
        lines = user.split("\n")
        return {"items": [{"n": i, "real": "оплата" not in t, "key": "чек клиенту" if i == 1 else "выдаём ли чек"}
                          for i, t in enumerate(lines, 1)]}


def test_conflicts_xlsx(tmp_path):
    pytest.importorskip("openpyxl")   # скрипт только для ПК, в образе сервиса openpyxl нет
    spec =importlib.util.spec_from_file_location("kbx", ROOT / "scripts" / "kb_conflicts_xlsx.py")
    kbx = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(kbx)
    (tmp_path / "rule_conflicts.json").write_text(
        '[{"topic": "оплата через Авито", "chosen": "не через Авито", "off": ["через Авито можно"]}]',
        encoding="utf-8")
    (tmp_path / "conflicts_core.json").write_text(json.dumps([
        {"section": "Game Pass · FAQ", "topic": "чек", "rules": "Пусто", "premium": "не выдаём",
         "gamenonstop": "по запросу", "chosen": "не выдаём"},
        {"section": "Game Pass · FAQ", "topic": "оплата", "rules": "переводом", "premium": "",
         "gamenonstop": "переводом на карту", "chosen": "переводом"},
        {"section": "Game Pass · FAQ", "topic": "одно", "rules": "не упоминается", "premium": "да",
         "gamenonstop": ""}], ensure_ascii=False), encoding="utf-8")
    (tmp_path / "conflicts_games.json").write_text(json.dumps([
        {"section": "Игры · FAQ", "topic": "чек?", "rules": "", "premium": "не выдаём",
         "gamenonstop": "выдаём по запросу на почту", "chosen": "не выдаём"}], ensure_ascii=False), encoding="utf-8")
    rules, rows = kbx.load(tmp_path)
    assert [r["topic"] for r in rows] == ["чек", "оплата", "чек?"]      # «одно» — второй стороны нет
    assert rows[0]["rules"] == "" and rows[0]["kind"] == kbx.PM_GNS and rows[1]["kind"] == kbx.RULE
    picked = asyncio.run(kbx.select(ReviewLLM(), rows))
    assert len(picked) == 1                                               # «оплата» отсеяна, «чек» склеен
    assert picked[0]["gamenonstop"] == "выдаём по запросу на почту"       # из двух — самая полная
    assert picked[0]["section"] == "Game Pass · FAQ; Игры · FAQ"
    kbx.build(rules, picked, tmp_path / "out.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "out.xlsx")
    ws = wb["Противоречия"]
    assert [c.value for c in ws[1]][4:7] == ["Правила", "PREMIUM MARKET", "GameNonStop"]
    assert [c.value for c in ws[2]][:8] == [1, kbx.PM_GNS, "Game Pass · FAQ; Игры · FAQ", "чек?", None, "не выдаём",
                                            "выдаём по запросу на почту", "не выдаём"]
    assert [c.value for c in wb["Правила между собой"][2]][:4] == [1, "оплата через Авито", "не через Авито",
                                                                  "через Авито можно"]


def test_dialog_reloads_kb_after_nightly(monkeypatch):
    """Ночная пересборка доходит до бота без перезапуска (раньше база читалась только при старте)."""
    from services.dialog import brain, worker
    stamp, loads = ["v1"], []
    monkeypatch.setattr(brain, "kb_stamp", lambda: stamp[0])
    monkeypatch.setattr(brain, "load_kb", lambda cls="core": loads.append(cls) or f"{stamp[0]}:{cls}")
    d = worker.Dialog.__new__(worker.Dialog)
    d.kb_at = None
    d.load_kb()
    d.load_kb()
    assert "v1:core" in d.system and loads.count("core") == 1
    stamp[0] = "v2"
    d.load_kb()
    assert "v2:core" in d.system and loads.count("core") == 2


def test_salvage_cut_conflicts():
    whole = '{"conflicts": [{"topic": "чек", "premium": "нет"}, {"topic": "PS Plus", "gamenonstop": "есть"}]}'
    assert [c["topic"] for c in build_kb.salvage(whole)] == ["чек", "PS Plus"]
    cut = whole[:whole.index("PS Plus") + 3]      # ответ оборвался на лимите токенов
    assert [c["topic"] for c in build_kb.salvage(cut)] == ["чек"]
    assert build_kb.salvage("не json") == [] and build_kb.salvage('{"conflicts": null}') == []


def test_drop_steam():
    """30.10: Steam не продаём — из базы знаний кодом вырезаются пункты и строки про Steam."""
    text = ("- Клиент просит ключ без надписей для VPN, ключ для Steam, но ключа нет, только гифт.\n"
            "- О товаре для платформы, с которой не работаем (Steam, PS, Xbox 360, GOG), пополнить баланс Steam.\n"
            "- Вопросы по другим устройствам (Steam Deck).\n"
            "- Смена региона стима — 24 часа.\n"
            "- Стимул купить подлиннее — подарок.")
    assert build_kb.drop_steam(text).split("\n") == [
        "- Клиент просит ключ без надписей для VPN, но ключа нет, только гифт.",
        "- О товаре для платформы, с которой не работаем (PS, Xbox 360, GOG).",
        "- Вопросы по другим устройствам (Steam Deck).",   # Steam Deck — ПК, не Steam
        "- Стимул купить подлиннее — подарок."]


class EnforceLLM:
    """Второй проход сборки (31.10): правило 1 расходится с фразой куска — правим; фраза, которая сама есть
    правило 2, — спор правил, не правим."""
    def __init__(self):
        self.fixed, self.usage = [], {}

    async def chat_json(self, system, user, **kw):
        if system == build_kb.kb_rules.PICK:
            return {"ids": [1, 2], "home": 1, "words": []}
        if system == build_kb.FIND:
            if "ПРАВИЛО: Оплата только заранее" in user:
                return {"conflicts": [{"id": 1, "quote": "Можем подключить, а потом оплатите."},
                                      {"id": 2, "quote": "Сначала войдите в наш аккаунт."}]}
            return {"conflicts": []}
        if system == build_kb.SURE:
            return {"real": True, "other": 2 if "Сначала войдите" in user.split("ФРАЗА:")[1][:40] else None}
        if system == build_kb.FIX:
            data = json.loads(user)
            self.fixed.append(data["phrases"])
            return {"text": "#### core › Частые вопросы\n" + data["text"].replace(
                "Можем подключить, а потом оплатите.", "Оплата только полная и заранее.")}
        raise AssertionError(system[:40])


def test_enforce_rule_wins_and_rule_clash(tmp_path):
    (tmp_path / "kb_core.md").write_text(
        "# База\n\n## Частые вопросы\n\n#### Можно оплатить после?\nМожем подключить, а потом оплатите.\n\n"
        "#### Как войти\nСначала войдите в наш аккаунт.\n", encoding="utf-8")
    (tmp_path / "kb_games.md").write_text("# Игры\n\n## Частые вопросы\n\nПро игры.\n", encoding="utf-8")
    llm = EnforceLLM()
    owner = ["Оплата только заранее, после подключения не оплачивают.", "Сначала войдите в наш аккаунт."]
    n, clash = asyncio.run(build_kb.enforce(llm, owner, tmp_path))
    text = (tmp_path / "kb_core.md").read_text(encoding="utf-8")
    assert "Оплата только полная и заранее." in text and "потом оплатите" not in text   # правило важнее переписки
    assert "Сначала войдите в наш аккаунт." in text                                       # спор правил — не трогаем
    assert "core ›" not in text                       # путь куска, который LLM вернула заголовком, вырезан
    assert llm.fixed == [["Можем подключить, а потом оплатите."]] and n == 1
    assert clash == [tuple(sorted(owner))]
